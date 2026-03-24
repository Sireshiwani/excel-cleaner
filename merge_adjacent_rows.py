#!/usr/bin/env python3
"""Merge data from each pair of adjacent rows in an Excel worksheet.

The header row (row 1) is preserved as-is. Starting from row 2, data rows are
processed in pairs (2+3, 4+5, ...):
  - If one value is empty, the non-empty value is kept.
  - If both values are equal, the value is kept once.
  - If both values are non-empty and different, they are concatenated.
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any, Iterable, List, Optional

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


def is_empty(value: Any) -> bool:
    """Return True for blank-like worksheet values."""
    return value is None or (isinstance(value, str) and value.strip() == "")


def merge_cell_values(first: Any, second: Any, separator: str = " ") -> Any:
    """Merge two cell values using simple, predictable rules."""
    if is_empty(first) and is_empty(second):
        return None
    if is_empty(first):
        return second
    if is_empty(second):
        return first
    if first == second:
        return first
    return f"{first}{separator}{second}"


def get_row_values(ws: Worksheet, row_number: int, max_col: int) -> List[Any]:
    """Read a worksheet row into a list of values."""
    return [ws.cell(row=row_number, column=col).value for col in range(1, max_col + 1)]


def merge_two_rows(first_row: Iterable[Any], second_row: Iterable[Any], separator: str) -> List[Any]:
    """Merge two row iterables column by column."""
    return [merge_cell_values(first, second, separator) for first, second in zip(first_row, second_row)]


def merge_adjacent_rows_in_worksheet(ws: Worksheet, separator: str = " ") -> None:
    """Merge adjacent data rows in-place while preserving the header row."""
    original_max_row = ws.max_row
    max_col = ws.max_column

    if original_max_row <= 1:
        return

    merged_rows: List[List[Any]] = []
    row_idx = 2
    while row_idx <= original_max_row:
        first_row = get_row_values(ws, row_idx, max_col)

        if row_idx + 1 <= original_max_row:
            second_row = get_row_values(ws, row_idx + 1, max_col)
            merged_rows.append(merge_two_rows(first_row, second_row, separator))
            row_idx += 2
        else:
            merged_rows.append(first_row)
            row_idx += 1

    ws.delete_rows(2, original_max_row - 1)
    for row in merged_rows:
        ws.append(row)


def build_output_path(input_path: Path, output_path: Optional[Path]) -> Path:
    """Build a default output path when one isn't explicitly provided."""
    if output_path is not None:
        return output_path
    return input_path.with_name(f"{input_path.stem}_merged{input_path.suffix}")


def process_workbook(
    input_path: Path,
    output_path: Optional[Path] = None,
    sheet_name: Optional[str] = None,
    separator: str = " ",
) -> Path:
    """Load, transform, and save a workbook."""
    if not input_path.exists():
        raise FileNotFoundError(f"Input file does not exist: {input_path}")

    wb = load_workbook(filename=input_path)
    ws = wb[sheet_name] if sheet_name else wb.active
    merge_adjacent_rows_in_worksheet(ws, separator=separator)

    final_output_path = build_output_path(input_path, output_path)
    wb.save(final_output_path)
    return final_output_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Merge data from each pair of adjacent rows in an Excel worksheet."
    )
    parser.add_argument("input_file", type=Path, help="Path to input .xlsx file")
    parser.add_argument(
        "-o",
        "--output-file",
        type=Path,
        default=None,
        help="Path to output .xlsx file (default: <input>_merged.xlsx)",
    )
    parser.add_argument(
        "-s",
        "--sheet",
        type=str,
        default=None,
        help="Optional sheet name (default: active sheet)",
    )
    parser.add_argument(
        "--separator",
        type=str,
        default=" ",
        help="Separator used when both cells are non-empty and different (default: space)",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output = process_workbook(
        input_path=args.input_file,
        output_path=args.output_file,
        sheet_name=args.sheet,
        separator=args.separator,
    )
    print(f"Merged workbook written to: {output}")


if __name__ == "__main__":
    main()
