import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from merge_adjacent_rows import process_workbook


class MergeAdjacentRowsTests(unittest.TestCase):
    def test_merges_pairs_and_preserves_header(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            input_path = Path(tmp_dir) / "input.xlsx"
            output_path = Path(tmp_dir) / "output.xlsx"

            wb = Workbook()
            ws = wb.active
            ws.append(["Name", "City", "Score"])
            ws.append(["Alice", "London", 10])
            ws.append(["", "London", 20])
            ws.append(["Bob", "", None])
            ws.append(["Bobby", "Paris", 5])
            wb.save(input_path)

            process_workbook(input_path=input_path, output_path=output_path, separator=" | ")

            merged_wb = load_workbook(output_path)
            merged_ws = merged_wb.active

            self.assertEqual(merged_ws.max_row, 3)
            self.assertEqual([cell.value for cell in merged_ws[1]], ["Name", "City", "Score"])
            self.assertEqual([cell.value for cell in merged_ws[2]], ["Alice", "London", "10 | 20"])
            self.assertEqual([cell.value for cell in merged_ws[3]], ["Bob | Bobby", "Paris", 5])

    def test_handles_odd_number_of_data_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            input_path = Path(tmp_dir) / "input_odd.xlsx"
            output_path = Path(tmp_dir) / "output_odd.xlsx"

            wb = Workbook()
            ws = wb.active
            ws.append(["Col1", "Col2"])
            ws.append(["A", "B"])
            ws.append(["A2", "B2"])
            ws.append(["LAST", "ROW"])
            wb.save(input_path)

            process_workbook(input_path=input_path, output_path=output_path)

            merged_wb = load_workbook(output_path)
            merged_ws = merged_wb.active

            self.assertEqual(merged_ws.max_row, 3)
            self.assertEqual([cell.value for cell in merged_ws[2]], ["A A2", "B B2"])
            self.assertEqual([cell.value for cell in merged_ws[3]], ["LAST", "ROW"])


if __name__ == "__main__":
    unittest.main()
